"""
TPL Comment Extractor — Web Application
========================================
Streamlit front-end for generating TPL Comments Excel sheets.
All extraction logic lives in extractor/extractor.py.
This file contains ONLY UI and orchestration code.
"""

import streamlit as st
import os
import sys
import time
import uuid
import shutil
import tempfile
from pathlib import Path
from datetime import datetime

# ── Path setup ────────────────────────────────────────────────────────────────
ROOT = Path(__file__).parent
sys.path.insert(0, str(ROOT))
UPLOADS_DIR = ROOT / "uploads"
OUTPUTS_DIR = ROOT / "outputs"
UPLOADS_DIR.mkdir(exist_ok=True)
OUTPUTS_DIR.mkdir(exist_ok=True)

# ── Import extraction engine ──────────────────────────────────────────────────
try:
    from extractor import run_extraction, ExtractionResult
    ENGINE_AVAILABLE = True
except ImportError as e:
    ENGINE_AVAILABLE = False
    ENGINE_ERROR = str(e)

# ── Excel generation (pure formatting, not extraction) ────────────────────────
try:
    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side, GradientFill
    )
    from openpyxl.utils import get_column_letter
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False


def build_excel(rows: list, debug_rows: list, output_path: str):
    """
    Build the TPL Comments Excel file from extraction result rows.
    Matches the exact format of the hand-made TPL_Comments.xlsx.
    Also writes a Debug sheet with all raw annotation data.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TPL Comments"

    # ── Column widths ─────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 7    # Sr No
    ws.column_dimensions["B"].width = 12   # Submittal
    ws.column_dimensions["C"].width = 72   # Document
    ws.column_dimensions["D"].width = 62   # Customer Comments
    ws.column_dimensions["E"].width = 24   # Xylem Remarks

    # ── Header row ────────────────────────────────────────────────────────────
    headers = ["Sr No.", "Submittal", "Document", "Customer Comments", "Xylem Remarks"]
    header_fill   = PatternFill("solid", fgColor="1E293B")  # dark steel
    header_font   = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    header_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font    = header_font
        cell.fill    = header_fill
        cell.alignment = header_align
        cell.border  = border

    ws.row_dimensions[1].height = 28

    # ── Data rows ─────────────────────────────────────────────────────────────
    alt_fill    = PatternFill("solid", fgColor="F8F7F4")
    white_fill  = PatternFill("solid", fgColor="FFFFFF")
    cnr_fill    = PatternFill("solid", fgColor="FEF3C7")   # amber for "Comment not Received"
    cmt_fill    = PatternFill("solid", fgColor="E6F4ED")   # green for rows with comments

    base_font   = Font(name="Arial", size=9)
    sr_font     = Font(name="Arial", size=9, color="6B7280")
    sub_font    = Font(name="Arial", size=9, bold=True)
    cmt_font    = Font(name="Arial", size=9, color="0D6E3F")
    cnr_font    = Font(name="Arial", size=9, color="B45309", italic=True)

    left_align  = Alignment(horizontal="left",   vertical="top", wrap_text=True)
    center_align= Alignment(horizontal="center", vertical="top", wrap_text=False)

    for row_idx, row in enumerate(rows, 2):
        has_comment = bool(row.get("customer_comments", "").strip())
        row_fill    = cmt_fill if has_comment else (alt_fill if row_idx % 2 == 0 else white_fill)

        # Sr No
        c = ws.cell(row=row_idx, column=1, value=row.get("sr_no", row_idx - 1))
        c.font = sr_font; c.fill = row_fill; c.alignment = center_align; c.border = border

        # Submittal
        c = ws.cell(row=row_idx, column=2, value=row.get("submittal", ""))
        c.font = sub_font; c.fill = row_fill; c.alignment = left_align; c.border = border

        # Document
        c = ws.cell(row=row_idx, column=3, value=row.get("document", ""))
        c.font = base_font; c.fill = row_fill; c.alignment = left_align; c.border = border

        # Customer Comments
        comments = row.get("customer_comments", "")
        c = ws.cell(row=row_idx, column=4, value=comments)
        c.font = cmt_font if comments else base_font
        c.fill = row_fill; c.alignment = left_align; c.border = border

        # Xylem Remarks
        remarks = row.get("xylem_remarks", "")
        c = ws.cell(row=row_idx, column=5, value=remarks)
        c.font = cnr_font if remarks == "Comment not Received" else base_font
        c.fill = cnr_fill if remarks == "Comment not Received" else row_fill
        c.alignment = left_align; c.border = border

        # Auto row height for wrapped text (approx)
        doc = row.get("document", "") or ""
        cmt = row.get("customer_comments","") or ""
        doc_lines = len(str(doc).split("\n"))
        cmt_lines = max(1, len([l for l in str(cmt).split("\n") if l.strip()]))
        ws.row_dimensions[row_idx].height = max(18, max(doc_lines, cmt_lines) * 14)

    # ── Freeze top row ────────────────────────────────────────────────────────
    ws.freeze_panes = "A2"

    # ── Auto filter ───────────────────────────────────────────────────────────
    ws.auto_filter.ref = f"A1:E{len(rows)+1}"

    # ── Debug sheet — All Annotations ────────────────────────────────────────
    if debug_rows:
        wd = wb.create_sheet("Debug - All Annotations")
        debug_headers = [
            "Submittal", "File", "Page",
            "Annotation Type", "Author", "Content (raw)", "Classified As"
        ]
        hfill  = PatternFill("solid", fgColor="1E293B")
        hfont  = Font(name="Arial", bold=True, color="FFFFFF", size=9)
        halign = Alignment(horizontal="center", vertical="center")
        for ci, h in enumerate(debug_headers, 1):
            cell = wd.cell(row=1, column=ci, value=h)
            cell.font = hfont; cell.fill = hfill
            cell.alignment = halign; cell.border = border
        debug_col_widths = [12, 40, 6, 22, 20, 80, 18]
        for ci, w in enumerate(debug_col_widths, 1):
            wd.column_dimensions[get_column_letter(ci)].width = w
        act_fill = PatternFill("solid", fgColor="E6FFE6")
        for ri, d in enumerate(debug_rows, 2):
            classified = "✅ COMMENT" if d.get("actionable") else "⛔ Not a comment"
            vals = [
                d.get("sub", ""), d.get("file", ""), d.get("page", ""),
                d.get("type", ""), d.get("author", ""),
                d.get("content", "")[:300], classified,
            ]
            for ci, val in enumerate(vals, 1):
                cell = wd.cell(row=ri, column=ci, value=val)
                cell.font = Font(name="Arial", size=9)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border
                if d.get("actionable"):
                    cell.fill = act_fill
        wd.freeze_panes = "A2"

    wb.save(output_path)


# ── PASSWORD PROTECTION ────────────────────────────────────────────────────
PASSWORD = "admin123"  # Change this to your secure password!

def check_password():
    """Returns True if password is correct"""
    if "password_correct" not in st.session_state:
        st.session_state.password_correct = False

    if not st.session_state.password_correct:
        st.markdown("### 🔐 Login Required")
        with st.form("password_form"):
            password = st.text_input("Enter password:", type="password")
            submitted = st.form_submit_button("Login")
            
            if submitted:
                if password == PASSWORD:
                    st.session_state.password_correct = True
                    st.success("✅ Login successful!")
                    st.rerun()
                else:
                    st.error("❌ Incorrect password. Try again.")
        return False
    return True

# Check password before showing app
if not check_password():
    st.stop()

# ─────────────────────────────────────────────────────────────────────────────# ── Page config ───────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="TPL Comment Extractor",
    page_icon="ðﾟﾓﾄ",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# ── Custom CSS ────────────────────────────────────────────────────────────────
st.markdown("""
<style>
/* ── Base ── */
[data-testid="stAppViewContainer"] {
    background: #f5f4f0;
}
[data-testid="stHeader"] { display: none; }

/* ── Main content width ── */
.main .block-container {
    max-width: 960px;
    padding: 0 2rem 4rem;
}

/* ── App header ── */
.tpl-header {
    background: #1e293b;
    color: white;
    padding: 20px 32px;
    margin: -1rem -2rem 2rem;
    border-bottom: 3px solid #1a56db;
    display: flex;
    align-items: center;
    gap: 16px;
}
.tpl-header .brand {
    font-family: 'Courier New', monospace;
    font-size: 13px;
    font-weight: 700;
    letter-spacing: 3px;
    text-transform: uppercase;
    color: #94a3b8;
}
.tpl-header h1 {
    font-size: 18px !important;
    font-weight: 700 !important;
    color: white !important;
    margin: 0 !important;
    padding: 0 !important;
}
.tpl-header .sub {
    font-size: 12px;
    color: #64748b;
    margin-top: 2px;
}

/* ── Section cards ── */
.tpl-card {
    background: white;
    border: 1px solid #e8e6e0;
    border-radius: 8px;
    padding: 24px 28px;
    margin-bottom: 20px;
}
.tpl-section-label {
    font-size: 10px;
    font-weight: 700;
    text-transform: uppercase;
    letter-spacing: 1.5px;
    color: #6b7280;
    margin-bottom: 14px;
    padding-bottom: 8px;
    border-bottom: 1px solid #e8e6e0;
    display: flex;
    align-items: center;
    gap: 8px;
}
.tpl-section-label .dot {
    width: 6px; height: 6px;
    background: #1a56db;
    border-radius: 50%;
    display: inline-block;
}

/* ── File chips ── */
.file-chip {
    display: inline-flex;
    align-items: center;
    gap: 6px;
    background: #f5f4f0;
    border: 1px solid #e8e6e0;
    border-radius: 4px;
    padding: 5px 10px;
    font-size: 12px;
    font-family: 'Courier New', monospace;
    margin: 3px;
}
.file-chip .ok { color: #0d6e3f; font-weight: 700; }

/* ── Stats ── */
.stats-row {
    display: flex;
    gap: 16px;
    margin-top: 16px;
}
.stat-box {
    flex: 1;
    background: #f5f4f0;
    border: 1px solid #e8e6e0;
    border-radius: 6px;
    padding: 16px;
    text-align: center;
}
.stat-box .num {
    font-size: 28px;
    font-weight: 700;
    font-family: 'Courier New', monospace;
    color: #1a56db;
}
.stat-box .lbl {
    font-size: 10px;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    color: #6b7280;
    margin-top: 2px;
}
.stat-box.green .num { color: #0d6e3f; }
.stat-box.amber .num { color: #b45309; }
.stat-box.red   .num { color: #c0392b; }

/* ── Info note ── */
.info-note {
    background: #e8f0fe;
    border-left: 3px solid #1a56db;
    border-radius: 0 4px 4px 0;
    padding: 10px 14px;
    font-size: 12px;
    color: #1e40af;
    margin-bottom: 16px;
    line-height: 1.6;
}

/* ── Error note ── */
.error-note {
    background: #fdecea;
    border-left: 3px solid #c0392b;
    border-radius: 0 4px 4px 0;
    padding: 10px 14px;
    font-size: 12px;
    color: #c0392b;
    margin-bottom: 10px;
    line-height: 1.6;
}

/* ── Buttons ── */
.stButton > button {
    font-family: 'Inter', sans-serif !important;
    font-weight: 600 !important;
    border-radius: 4px !important;
    transition: opacity 0.15s !important;
}
.stButton > button:hover { opacity: 0.88 !important; }

/* ── Download button ── */
.stDownloadButton > button {
    background: #0d6e3f !important;
    color: white !important;
    font-weight: 700 !important;
    border: none !important;
    padding: 12px 24px !important;
    font-size: 14px !important;
    border-radius: 4px !important;
    width: 100%;
}

/* ── Log box ── */
.log-box {
    background: #0d1117;
    border-radius: 4px;
    padding: 12px 14px;
    font-family: 'Courier New', monospace;
    font-size: 11px;
    color: #94a3b8;
    max-height: 180px;
    overflow-y: auto;
    line-height: 1.9;
    margin-top: 12px;
}

/* ── Hide streamlit branding ── */
#MainMenu, footer, [data-testid="stToolbar"] { visibility: hidden; }
</style>
""", unsafe_allow_html=True)


# ── Header ────────────────────────────────────────────────────────────────────
st.markdown("""
<div class="tpl-header">
  <div>
    <div class="brand">TPL · Dholera Fab 1</div>
    <h1>Generate TPL Comment Sheets Automatically</h1>
    <div class="sub">Upload ZIP packages → Extract Comments → Download Excel</div>
  </div>
</div>
""", unsafe_allow_html=True)


# ── Engine check ──────────────────────────────────────────────────────────────
if not ENGINE_AVAILABLE:
    st.markdown(f"""
    <div class="error-note">
        <strong>⚠ Extraction engine not found.</strong><br/>
        Make sure <code>extractor/extractor.py</code> exists in the project folder.<br/>
        Error: {ENGINE_ERROR}
    </div>
    """, unsafe_allow_html=True)
    st.stop()

if not EXCEL_AVAILABLE:
    st.markdown("""
    <div class="error-note">
        <strong>⚠ openpyxl not installed.</strong><br/>
        Run: <code>pip install openpyxl</code>
    </div>
    """, unsafe_allow_html=True)
    st.stop()


# ── Session state ─────────────────────────────────────────────────────────────
if "result"       not in st.session_state: st.session_state.result       = None
if "excel_bytes"  not in st.session_state: st.session_state.excel_bytes  = None
if "excel_name"   not in st.session_state: st.session_state.excel_name   = None
if "elapsed"      not in st.session_state: st.session_state.elapsed      = None
if "log_lines"    not in st.session_state: st.session_state.log_lines    = []


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 1 — UPLOAD
# ═══════════════════════════════════════════════════════════════════════════════
st.markdown('<div class="tpl-card">', unsafe_allow_html=True)
st.markdown('<div class="tpl-section-label"><span class="dot"></span>Step 1 — Upload ZIP Files</div>', unsafe_allow_html=True)

st.markdown("""
<div class="info-note">
  Upload the ZIP packages received from the submittal workflow.<br/>
  Each ZIP file = one submittal row in the output Excel.
  You can select multiple ZIPs at once (hold Ctrl or Cmd).
</div>
""", unsafe_allow_html=True)

uploaded_files = st.file_uploader(
    label="Drop ZIP files here or click to browse",
    type=["zip"],
    accept_multiple_files=True,
    label_visibility="collapsed",
    key="file_uploader",
)
st.markdown('</div>', unsafe_allow_html=True)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 2 — UPLOADED FILES LIST
# ═══════════════════════════════════════════════════════════════════════════════
if uploaded_files:
    st.markdown('<div class="tpl-card">', unsafe_allow_html=True)
    st.markdown('<div class="tpl-section-label"><span class="dot"></span>Step 2 — Uploaded Files</div>', unsafe_allow_html=True)

    chips_html = ""
    for f in uploaded_files:
        size_kb = f.size // 1024
        chips_html += f'<div class="file-chip"><span class="ok">▪</span> {f.name} <span style="color:#9ca3af;font-size:10px;">({size_kb} KB)</span></div>'

    st.markdown(chips_html, unsafe_allow_html=True)
    st.markdown(f"<br/><strong style='font-size:13px;'>Total: {len(uploaded_files)} ZIP file{'s' if len(uploaded_files)!=1 else ''} ready to process</strong>", unsafe_allow_html=True)
    st.markdown('</div>', unsafe_allow_html=True)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 3 — GENERATE BUTTON + PROCESSING
# ═══════════════════════════════════════════════════════════════════════════════
st.markdown('<div class="tpl-card">', unsafe_allow_html=True)
st.markdown('<div class="tpl-section-label"><span class="dot"></span>Step 3 — Generate Excel</div>', unsafe_allow_html=True)

col_btn, col_tip = st.columns([1, 2])

with col_btn:
    generate_clicked = st.button(
        "▶  Generate Excel",
        type="primary",
        disabled=not bool(uploaded_files),
        use_container_width=True,
    )

with col_tip:
    if not uploaded_files:
        st.markdown("<span style='font-size:12px;color:#6b7280;'>Upload ZIP files first to enable processing.</span>", unsafe_allow_html=True)
    else:
        st.markdown(f"<span style='font-size:12px;color:#6b7280;'>Ready to process <strong>{len(uploaded_files)} submittal{'s' if len(uploaded_files)!=1 else ''}</strong>. Click Generate to start.</span>", unsafe_allow_html=True)


# ── Processing ────────────────────────────────────────────────────────────────
if generate_clicked and uploaded_files:
    # Reset previous result
    st.session_state.result      = None
    st.session_state.excel_bytes = None
    st.session_state.log_lines   = []

    # Save uploaded files to temp disk location
    session_id  = str(uuid.uuid4())[:8]
    session_dir = UPLOADS_DIR / session_id
    session_dir.mkdir(parents=True, exist_ok=True)

    zip_paths = []
    for uf in uploaded_files:
        dest = session_dir / uf.name
        dest.write_bytes(uf.read())
        zip_paths.append(str(dest))

    # UI elements for progress
    status_text = st.empty()
    progress_bar = st.progress(0)
    log_placeholder = st.empty()
    log_lines = []

    def progress_cb(current: int, total: int, message: str):
        pct = int((current / total) * 100) if total > 0 else 0
        progress_bar.progress(pct)
        status_text.markdown(f"<span style='font-size:13px;color:#1a56db;font-weight:600;'>{message}</span>", unsafe_allow_html=True)
        log_lines.append(message)
        log_html = "".join(f"<div>{ln}</div>" for ln in log_lines[-12:])
        log_placeholder.markdown(f'<div class="log-box">{log_html}</div>', unsafe_allow_html=True)

    # ── Run extraction ────────────────────────────────────────────────────────
    t_start = time.time()
    try:
        result = run_extraction(
            zip_paths=zip_paths,
            progress_cb=progress_cb,
        )
        elapsed = round(time.time() - t_start, 1)
        st.session_state.result  = result
        st.session_state.elapsed = elapsed
        progress_bar.progress(100)
        status_text.markdown("<span style='font-size:13px;color:#0d6e3f;font-weight:600;'>✓ Extraction complete</span>", unsafe_allow_html=True)

    except Exception as e:
        st.markdown(f'<div class="error-note"><strong>Extraction failed:</strong> {str(e)}</div>', unsafe_allow_html=True)
        st.stop()

    # ── Build Excel ───────────────────────────────────────────────────────────
    try:
        ts         = datetime.now().strftime("%Y%m%d_%H%M%S")
        excel_name = f"TPL_Comments_{ts}.xlsx"
        excel_path = str(OUTPUTS_DIR / excel_name)
        build_excel(result.rows, getattr(result, 'debug_rows', []), excel_path)
        with open(excel_path, "rb") as f:
            st.session_state.excel_bytes = f.read()
        st.session_state.excel_name = excel_name
    except Exception as e:
        st.markdown(f'<div class="error-note"><strong>Excel generation failed:</strong> {str(e)}</div>', unsafe_allow_html=True)

    # ── Cleanup temp uploads ──────────────────────────────────────────────────
    try:
        shutil.rmtree(session_dir)
    except Exception:
        pass

st.markdown('</div>', unsafe_allow_html=True)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 4 — RESULTS
# ═══════════════════════════════════════════════════════════════════════════════
result = st.session_state.result

if result is not None:
    st.markdown('<div class="tpl-card">', unsafe_allow_html=True)
    st.markdown('<div class="tpl-section-label"><span class="dot"></span>Step 4 — Results</div>', unsafe_allow_html=True)

    elapsed_str = f"{st.session_state.elapsed}s" if st.session_state.elapsed else "—"

    st.markdown(f"""
    <div class="stats-row">
        <div class="stat-box">
            <div class="num">{result.total}</div>
            <div class="lbl">Submittals Processed</div>
        </div>
        <div class="stat-box green">
            <div class="num">{result.with_comments}</div>
            <div class="lbl">With Comments</div>
        </div>
        <div class="stat-box amber">
            <div class="num">{result.no_comments}</div>
            <div class="lbl">Comment not Received</div>
        </div>
        <div class="stat-box">
            <div class="num">{elapsed_str}</div>
            <div class="lbl">Processing Time</div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    # Non-fatal errors
    if result.errors:
        st.markdown("<br/>", unsafe_allow_html=True)
        for err in result.errors:
            st.markdown(f'<div class="error-note">⚠ {err}</div>', unsafe_allow_html=True)

    # Preview table
    if result.rows:
        st.markdown("<br/><strong style='font-size:13px;'>Preview</strong>", unsafe_allow_html=True)

        import pandas as pd
        preview_data = []
        for row in result.rows:
            doc = row.get("document", "") or ""
            cmt = row.get("customer_comments","") or ""
            doc_short = doc[:80] + "_" if doc and len(doc) > 80 else (doc or "")
            cmt_short = cmt[:60] + "…" if cmt and len(cmt) > 60 else (cmt or "")
            preview_data.append({
                "Sr No": row.get("sr_no",""),
                "Submittal": row.get("submittal",""),
                "Document": doc_short.replace("\n"," | "),
                "Comments": cmt_short or "—",
                "Xylem Remarks": row.get("xylem_remarks",""),
            })

        df = pd.DataFrame(preview_data)
        st.dataframe(
            df,
            use_container_width=True,
            hide_index=True,
            column_config={
                "Sr No":        st.column_config.NumberColumn(width="small"),
                "Submittal":    st.column_config.TextColumn(width="small"),
                "Document":     st.column_config.TextColumn(width="large"),
                "Comments":     st.column_config.TextColumn(width="large"),
                "Xylem Remarks":st.column_config.TextColumn(width="medium"),
            }
        )

    st.markdown('</div>', unsafe_allow_html=True)


# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 5 — DOWNLOAD
# ═══════════════════════════════════════════════════════════════════════════════
if st.session_state.excel_bytes:
    st.markdown('<div class="tpl-card">', unsafe_allow_html=True)
    st.markdown('<div class="tpl-section-label"><span class="dot"></span>Step 5 — Download</div>', unsafe_allow_html=True)

    st.markdown("""
    <div class="info-note">
      Your Excel file is ready. Click below to download.<br/>
      <strong>Review the file before sharing</strong> — the extracted comments are a draft for human review.
    </div>
    """, unsafe_allow_html=True)

    st.download_button(
        label=f"⬇  Download {st.session_state.excel_name}",
        data=st.session_state.excel_bytes,
        file_name=st.session_state.excel_name,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )
    st.markdown('</div>', unsafe_allow_html=True)


# ═══════════════════════════════════════════════════════════════════════════════
# FOOTER
# ═══════════════════════════════════════════════════════════════════════════════
st.markdown("""
<div style="text-align:center;padding:32px 0 8px;font-size:11px;color:#9ca3af;">
  TPL Comment Extractor · Internal Tool · Dholera Fab 1 Project<br/>
  Generated Excel is a draft for human review — verify before submission.
</div>
""", unsafe_allow_html=True)
